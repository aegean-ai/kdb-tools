"""Spreadsheet cells as evidence: read a cell, find the cells that hold a value, compare values.

A cell source cites a value by file, sheet and A1 cell, such as sheet "DH", cell "B4". .xlsx
and .xlsm files are read with openpyxl, using the value each formula last computed (macros never run); .csv files have one
unnamed sheet, so their sources name no sheet.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

SUFFIXES = (".xlsx", ".xlsm", ".csv")
A1 = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")


class SheetError(ValueError):
    """The sheet or cell cannot be read; the message says why."""


def needs_sheet(path: Path) -> bool:
    return path.suffix.lower() in (".xlsx", ".xlsm")


def is_a1(ref) -> bool:
    return isinstance(ref, str) and bool(A1.match(ref))


def column_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def column_letters(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def same_value(cited, held) -> bool:
    """Numbers match by value (169.77 and "169.770"); anything else by its text, ignoring outer spaces."""
    try:
        a, b = float(str(cited).strip()), float(str(held).strip())
    except ValueError:
        return str(cited).strip() == str(held).strip()
    return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))


def show(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


class Workbook:
    """One spreadsheet file, read once: {sheet name: {(row, column): value}}, plus the formula cells with no saved value."""

    def __init__(self, path: Path):
        self.path = path
        self.sheets: dict[str, dict[tuple[int, int], object]] = {}
        self.unsaved: set[tuple[str, int, int]] = set()
        if path.suffix.lower() == ".csv":
            with open(path, newline="", encoding="utf-8-sig") as fh:
                self.sheets[""] = {(r, c): v for r, row in enumerate(csv.reader(fh), 1) for c, v in enumerate(row, 1) if v != ""}
            return
        from openpyxl import load_workbook

        values = load_workbook(path, data_only=True, read_only=True)
        formulas = load_workbook(path, data_only=False, read_only=True)
        try:
            for ws in values.worksheets:
                cells = {}
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            cells[(cell.row, cell.column)] = cell.value
                for row in formulas[ws.title].iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("=") and (cell.row, cell.column) not in cells:
                            self.unsaved.add((ws.title, cell.row, cell.column))
                self.sheets[ws.title] = cells
        finally:
            values.close()
            formulas.close()

    def _sheet(self, sheet: str | None) -> str:
        if self.path.suffix.lower() == ".csv":
            return ""
        if sheet not in self.sheets:
            raise SheetError(f"there is no sheet {sheet!r}; the sheets are: {', '.join(self.sheets)}")
        return sheet

    def read(self, sheet: str | None, ref: str):
        """The value in a cell, or None when it is empty."""
        name = self._sheet(sheet)
        m = A1.match(ref)
        if not m:
            raise SheetError(f"{ref!r} is not an A1 cell reference such as B4")
        row, col = int(m.group(2)), column_index(m.group(1))
        if (name, row, col) in self.unsaved:
            raise SheetError(f"{ref} holds a formula with no saved value; open the file in a spreadsheet program and save it")
        return self.sheets[name].get((row, col))

    def find(self, value, sheet: str | None = None) -> list[str]:
        """Every cell that holds the value, as "Sheet!B4" (or "B4" in a .csv), in sheet then row then column order."""
        names = [self._sheet(sheet)] if sheet is not None or self.path.suffix.lower() == ".csv" else list(self.sheets)
        out = []
        for name in names:
            for (r, c), held in sorted(self.sheets[name].items()):
                if same_value(value, held):
                    ref = f"{column_letters(c)}{r}"
                    out.append(f"{name}!{ref}" if name else ref)
        return out
